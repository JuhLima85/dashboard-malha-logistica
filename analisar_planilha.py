from openpyxl import load_workbook

ARQUIVO = "base_completa.xlsx"
ABA = "Base"

print("\n========================================")
print("   COLUNAS DA BASE COMPLETA")
print("========================================\n")

wb = load_workbook(
    ARQUIVO,
    read_only=True,
    data_only=True
)

ws = wb[ABA]

cabecalho = next(
    ws.iter_rows(
        min_row=1,
        max_row=1,
        values_only=True
    )
)

for numero, coluna in enumerate(cabecalho, start=1):
    if coluna is not None:
        print(f"{numero:02d} - {coluna}")

print("\n========================================")
print(f"TOTAL DE COLUNAS: {len(cabecalho)}")
print("========================================")

wb.close()